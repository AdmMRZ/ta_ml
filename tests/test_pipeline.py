from __future__ import annotations

import hashlib
import os
import tempfile
import unittest
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

from ta_ml.cleaning import (
    RAW_GEOMETRY_FILENAME,
    RAW_LIGHTNING_FILENAME,
    RAW_STATIC_FILENAME,
    build_clean_workbook,
)

ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = Path(os.environ.get("TA_ML_DATA_DIR", ROOT / "data"))


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


class CleanWorkbookTests(unittest.TestCase):
    def test_builds_expected_population_without_changing_raw_sources(self) -> None:
        raw_paths = [
            DATA_DIR / RAW_LIGHTNING_FILENAME,
            DATA_DIR / RAW_STATIC_FILENAME,
            DATA_DIR / RAW_GEOMETRY_FILENAME,
        ]
        before = {path: sha256(path) for path in raw_paths}

        with tempfile.TemporaryDirectory() as temporary_directory:
            output_path = Path(temporary_directory) / "ta_ml_clean_v2.xlsx"
            summary = build_clean_workbook(
                source_dir=DATA_DIR,
                output_path=output_path,
            )

            self.assertTrue(output_path.is_file())
            self.assertEqual(summary.tower_year_rows, 3702)
            self.assertEqual(summary.unique_towers, 617)
            self.assertEqual(summary.static_snapshot_rows, 617)

            workbook = load_workbook(output_path, read_only=True, data_only=True)
            try:
                self.assertEqual(
                    workbook.sheetnames,
                    [
                        "README",
                        "tower_year",
                        "tower_registry",
                        "coordinate_2020_audit",
                        "source_notes",
                    ],
                )
                tower_year = workbook["tower_year"]
                header = next(tower_year.iter_rows(min_row=1, max_row=1, values_only=True))
                rows = list(tower_year.iter_rows(min_row=2, values_only=True))
                self.assertEqual(len(rows), 3702)
                column = {name: index for index, name in enumerate(header)}
                keys = {
                    (row[column["corridor"]], row[column["tower_key"]], row[column["year"]])
                    for row in rows
                }
                self.assertEqual(len(keys), 3702)
                self.assertEqual(
                    Counter(row[column["corridor"]] for row in rows),
                    {"PLK-KSG": 1164, "KSG-SKS": 1416, "KSG-SDN": 1122},
                )
                self.assertEqual(
                    Counter(
                        (row[column["year"]], row[column["coordinate_registry_match"]])
                        for row in rows
                    ),
                    {
                        (2020, True): 609,
                        (2020, None): 8,
                        **{(year, True): 617 for year in range(2021, 2026)},
                    },
                )
                self.assertEqual(
                    Counter(row[column["reconciliation_status"]] for row in rows),
                    {
                        "exact_coordinate_reassigned": 609,
                        "unresolved_2020_no_count": 8,
                        "source_name_coordinate_verified": 3085,
                    },
                )
                self.assertEqual(
                    sum(
                        row[column["count"]] is None for row in rows if row[column["year"]] == 2020
                    ),
                    8,
                )
                corrected = next(
                    row
                    for row in rows
                    if row[column["year"]] == 2020
                    and row[column["tower_key"]] == "KASONGAN - SKS #233"
                )
                self.assertEqual(corrected[column["count"]], 11)
                self.assertEqual(corrected[column["source_tower_name"]], "KASONGAN - SUDAN #186")

                registry = workbook["tower_registry"]
                registry_header = next(registry.iter_rows(max_row=1, values_only=True))
                registry_rows = list(registry.iter_rows(min_row=2, values_only=True))
                self.assertEqual(len(registry_rows), 617)
                registry_column = {name: index for index, name in enumerate(registry_header)}
                self.assertEqual(
                    len(
                        {
                            (r[registry_column["corridor"]], r[registry_column["tower_key"]])
                            for r in registry_rows
                        }
                    ),
                    617,
                )
                for field in ("elevation_m", "tower_height_m", "span_left_m", "span_right_m"):
                    self.assertTrue(
                        all(
                            isinstance(r[registry_column[field]], (int, float))
                            for r in registry_rows
                        )
                    )
                self.assertTrue(
                    all(r[registry_column["geometry_source_row"]] for r in registry_rows)
                )
                self.assertTrue(all(r[registry_column["static_source_row"]] for r in registry_rows))
                self.assertEqual(
                    sum(
                        r[registry_column["grounding_d_matches_sheet4"]] is False
                        for r in registry_rows
                    ),
                    126,
                )
                self.assertEqual(
                    sum(
                        r[registry_column["grounding_total_matches_sheet4"]] is False
                        for r in registry_rows
                    ),
                    143,
                )

                audit = workbook["coordinate_2020_audit"]
                audit_header = next(audit.iter_rows(max_row=1, values_only=True))
                audit_rows = list(audit.iter_rows(min_row=2, values_only=True))
                audit_column = {name: index for index, name in enumerate(audit_header)}
                self.assertEqual(
                    Counter(r[audit_column["match_status"]] for r in audit_rows),
                    {"exact_coordinate": 609, "unmatched": 8},
                )
                example = next(
                    r
                    for r in audit_rows
                    if r[audit_column["source_tower_key"]] == "KASONGAN - SUDAN #186"
                )
                self.assertEqual(example[audit_column["matched_tower_key"]], "KASONGAN - SKS #233")
            finally:
                workbook.close()

        after = {path: sha256(path) for path in raw_paths}
        self.assertEqual(before, after)


if __name__ == "__main__":
    unittest.main()
