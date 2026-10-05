"""Select with forward-year validation, refit through 2024, and evaluate 2025."""

from __future__ import annotations

import json
import platform
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
import sklearn
from openpyxl import load_workbook
from sklearn.compose import ColumnTransformer
from sklearn.ensemble import HistGradientBoostingRegressor
from sklearn.impute import SimpleImputer
from sklearn.linear_model import PoissonRegressor
from sklearn.metrics import mean_absolute_error, mean_poisson_deviance
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import OneHotEncoder, StandardScaler

from .cleaning import DataValidationError, normalize_text
from .config import (
    COUNT_DENSITY_DIVISOR,
    DEVELOPMENT_END_YEAR,
    SELECTION_FOLDS,
    TARGET_YEARS,
    TEST_YEAR,
    TOWER_COUNT,
)


@dataclass(frozen=True)
class EvaluationResult:
    report_path: Path
    predictions_path: Path


HISTORY = ["count_lag1", "count_lag2", "count_mean_available", "count_history_years"]
LOCATION = ["latitude_verified", "longitude_verified"]
GEOMETRY = [
    "elevation_m",
    "tower_height_m",
    "span_left_m",
    "span_right_m",
    "avg_neighbor_distance_m",
    "terrain_factor",
    "angle_exposure_factor",
]
PROTECTION_ASOF = ["tla_active_asof"]
SNAPSHOT_GROUNDING = [
    "grounding_leg_a_ohm_snapshot_2026",
    "grounding_leg_b_ohm_snapshot_2026",
    "grounding_leg_c_ohm_snapshot_2026",
    "grounding_rod_count_snapshot_2026",
    "soil_resistivity_snapshot_2026",
    "protection_angle_deg_snapshot_2026",
]
FEATURE_GROUPS = {
    "history": HISTORY,
    "history_location": HISTORY + LOCATION,
    "history_location_geometry": HISTORY + LOCATION + GEOMETRY,
    "history_location_geometry_tla": HISTORY + LOCATION + GEOMETRY + PROTECTION_ASOF,
}
MODELS = ("poisson_regressor", "hist_gradient_boosting_poisson")


def _pipeline(model_name: str, numerical: list[str], categorical: list[str]) -> Pipeline:
    numeric_pipeline = Pipeline(
        [
            ("impute", SimpleImputer(strategy="median", keep_empty_features=True)),
            ("scale", StandardScaler()),
        ]
    )
    categorical_pipeline = Pipeline(
        [
            ("impute", SimpleImputer(strategy="most_frequent")),
            ("onehot", OneHotEncoder(handle_unknown="ignore", sparse_output=False)),
        ]
    )
    transform = ColumnTransformer(
        [
            ("numeric", numeric_pipeline, numerical),
            ("categorical", categorical_pipeline, categorical),
        ]
    )
    if model_name == "poisson_regressor":
        regressor = PoissonRegressor(alpha=1.0, max_iter=1000)
    elif model_name == "hist_gradient_boosting_poisson":
        regressor = HistGradientBoostingRegressor(
            loss="poisson",
            max_iter=150,
            learning_rate=0.05,
            max_leaf_nodes=15,
            min_samples_leaf=20,
            random_state=42,
        )
    else:
        raise ValueError(f"Unknown model: {model_name}")
    return Pipeline([("prepare", transform), ("model", regressor)])


def _metrics(actual: pd.Series | np.ndarray, predicted: np.ndarray) -> dict[str, float]:
    y_pred = np.asarray(predicted, dtype=float)
    if not np.all(np.isfinite(y_pred)):
        raise DataValidationError("Model produced a non-finite prediction")
    return {
        "mae": float(mean_absolute_error(actual, y_pred)),
        "poisson_deviance": float(mean_poisson_deviance(actual, np.maximum(y_pred, 1e-9))),
    }


def _prepare_frame(frame: pd.DataFrame, numerical: list[str]) -> pd.DataFrame:
    result = frame[[*numerical, "corridor"]].copy()
    for column in numerical:
        result[column] = pd.to_numeric(result[column], errors="coerce")
    return result


def _read_pln(path: Path) -> pd.DataFrame:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook["CALC_ENGINE_R22"]
        iterator = worksheet.iter_rows(values_only=True)
        headers = next(iterator)
        index = {name: i for i, name in enumerate(headers)}
        required = {
            "Section",
            "Tower_Name",
            "Ng",
            "Collection_Area_m2",
            "Combined_Exposure_Factor",
            "P_strike_annual",
        }
        if not required.issubset(index):
            raise DataValidationError("PLN comparator columns are missing")
        records = [
            {
                "corridor": normalize_text(row[index["Section"]]),
                "tower_key": normalize_text(row[index["Tower_Name"]]),
                "pln_ng": row[index["Ng"]],
                "pln_collection_area_m2": row[index["Collection_Area_m2"]],
                "pln_exposure_factor": row[index["Combined_Exposure_Factor"]],
                "p_strike_annual_pln": row[index["P_strike_annual"]],
            }
            for row in iterator
            if row[index["Section"]] is not None
        ]
    finally:
        workbook.close()
    comparator = pd.DataFrame(records)
    for column in (
        "pln_ng",
        "pln_collection_area_m2",
        "pln_exposure_factor",
        "p_strike_annual_pln",
    ):
        comparator[column] = pd.to_numeric(comparator[column], errors="coerce")
    if len(comparator) != TOWER_COUNT or comparator.duplicated(["corridor", "tower_key"]).any():
        raise DataValidationError("PLN comparator must have 617 unique tower identities")
    if (
        comparator[
            ["pln_ng", "pln_collection_area_m2", "pln_exposure_factor", "p_strike_annual_pln"]
        ]
        .isna()
        .any()
        .any()
    ):
        raise DataValidationError("PLN exposure columns contain missing/non-numeric values")
    reconstructed = (
        comparator["pln_ng"]
        * comparator["pln_collection_area_m2"]
        / 1_000_000
        * comparator["pln_exposure_factor"]
    )
    if not np.allclose(reconstructed, comparator["p_strike_annual_pln"], atol=1e-9):
        raise DataValidationError("PLN P_strike_annual does not match its exposure inputs")
    return comparator


def evaluate_models(
    *, panel_path: Path, pln_path: Path, output_dir: Path, overwrite: bool = False
) -> EvaluationResult:
    """Select on 2023-2024; 2026-snapshot grounding is sensitivity-only."""

    panel_path, pln_path, output_dir = (
        panel_path.resolve(),
        pln_path.resolve(),
        output_dir.resolve(),
    )
    if not panel_path.is_file() or not pln_path.is_file():
        raise DataValidationError("Model panel or PLN comparator workbook is missing")
    report_path = output_dir / "evaluation_report.json"
    predictions_path = output_dir / "test_2025_predictions.csv"
    if not overwrite and (report_path.exists() or predictions_path.exists()):
        raise FileExistsError(f"Evaluation outputs already exist in {output_dir}")

    frame = pd.read_csv(panel_path, keep_default_na=True)
    required = set(HISTORY + LOCATION + GEOMETRY + PROTECTION_ASOF + SNAPSHOT_GROUNDING)
    required |= {"corridor", "tower_key", "split", "target_year", "target_count"}
    if not required.issubset(frame):
        raise DataValidationError(f"Missing model panel columns: {sorted(required - set(frame))}")
    if (
        len(frame) != TOWER_COUNT * len(TARGET_YEARS)
        or frame.duplicated(["corridor", "tower_key", "target_year"]).any()
    ):
        raise DataValidationError("Model panel population/keys are invalid")
    expected_split = {
        year: "development" if year <= DEVELOPMENT_END_YEAR else "test" for year in TARGET_YEARS
    }
    if any(row.split != expected_split.get(row.target_year) for row in frame.itertuples()):
        raise DataValidationError("Temporal split does not match the registered protocol")
    if frame["target_count"].isna().any() or (frame["target_count"] < 0).any():
        raise DataValidationError("Target Count is missing or negative")
    first_target = frame[frame["target_year"] == TARGET_YEARS[0]]
    matched_2020 = first_target["count_2020_coordinate_match_status"].eq(
        "exact_coordinate_reassigned"
    )
    unresolved_2020 = first_target["count_2020_coordinate_match_status"].eq("unmatched")
    if (
        len(first_target) != TOWER_COUNT
        or int(matched_2020.sum() + unresolved_2020.sum()) != TOWER_COUNT
        or first_target.loc[matched_2020, "count_lag1"].isna().any()
        or first_target.loc[unresolved_2020, "count_lag1"].notna().any()
    ):
        raise DataValidationError("2020 coordinate reconciliation and lag features disagree")

    train = frame[frame["target_year"] < DEVELOPMENT_END_YEAR].copy()
    validation = frame[frame["target_year"] == DEVELOPMENT_END_YEAR].copy()
    test = frame[frame["split"] == "test"].copy()
    if (len(train), len(validation), len(test)) != (
        TOWER_COUNT * (len(TARGET_YEARS) - 2),
        TOWER_COUNT,
        TOWER_COUNT,
    ):
        raise DataValidationError("Unexpected train/validation/test sizes")

    selection_folds = [
        {"train_target_years": list(years), "validation_target_year": year}
        for years, year in SELECTION_FOLDS
    ]
    candidates: list[dict[str, Any]] = []
    for group_name, numerical in FEATURE_GROUPS.items():
        for model_name in MODELS:
            fold_scores: list[dict[str, float | int]] = []
            for fold in selection_folds:
                fold_train = frame[frame["target_year"].isin(fold["train_target_years"])]
                fold_validation = frame[frame["target_year"] == fold["validation_target_year"]]
                model = _pipeline(model_name, numerical, ["corridor"])
                model.fit(_prepare_frame(fold_train, numerical), fold_train["target_count"])
                predictions = model.predict(_prepare_frame(fold_validation, numerical))
                fold_scores.append(
                    {
                        "validation_target_year": fold["validation_target_year"],
                        **_metrics(fold_validation["target_count"], predictions),
                    }
                )
            candidates.append(
                {
                    "feature_group": group_name,
                    "model": model_name,
                    "fold_scores": fold_scores,
                    "mean_mae": float(np.mean([score["mae"] for score in fold_scores])),
                    "mean_poisson_deviance": float(
                        np.mean([score["poisson_deviance"] for score in fold_scores])
                    ),
                }
            )
    selected = min(
        candidates,
        key=lambda item: (
            item["mean_mae"],
            item["mean_poisson_deviance"],
            len(FEATURE_GROUPS[item["feature_group"]]),
            MODELS.index(item["model"]),
        ),
    )
    selected_numerical = FEATURE_GROUPS[selected["feature_group"]]
    selected_model = _pipeline(selected["model"], selected_numerical, ["corridor"])
    final_train = frame[frame["target_year"] <= DEVELOPMENT_END_YEAR]
    selected_model.fit(_prepare_frame(final_train, selected_numerical), final_train["target_count"])

    # Diagnostic only: measurements in this 2026-labeled file lack as-of dates.
    sensitivity_features = selected_numerical + SNAPSHOT_GROUNDING
    sensitivity_model = _pipeline(selected["model"], sensitivity_features, ["corridor"])
    sensitivity_model.fit(_prepare_frame(train, sensitivity_features), train["target_count"])
    sensitivity_prediction = sensitivity_model.predict(
        _prepare_frame(validation, sensitivity_features)
    )
    sensitivity_metrics = _metrics(validation["target_count"], sensitivity_prediction)

    test_prediction = selected_model.predict(_prepare_frame(test, selected_numerical))
    test_metrics = _metrics(test["target_count"], test_prediction)
    baseline_validation = [
        {
            "validation_target_year": fold["validation_target_year"],
            **_metrics(
                frame.loc[frame["target_year"] == fold["validation_target_year"], "target_count"],
                frame.loc[
                    frame["target_year"] == fold["validation_target_year"], "count_lag1"
                ].to_numpy(dtype=float),
            ),
        }
        for fold in selection_folds
    ]
    baseline_test = _metrics(test["target_count"], test["count_lag1"].to_numpy(dtype=float))

    results = test[["corridor", "tower_key", "target_year", "target_count"]].copy()
    results["prediction_ml"] = test_prediction
    results["prediction_last_year"] = test["count_lag1"].to_numpy(dtype=float)
    results = results.merge(
        _read_pln(pln_path), on=["corridor", "tower_key"], validate="one_to_one"
    )
    if len(results) != TOWER_COUNT:
        raise DataValidationError("PLN comparator did not match all 617 test towers")
    results["observed_ng_2025"] = (results["target_count"] / COUNT_DENSITY_DIVISOR).round(2)
    results["prediction_ng_2025"] = (results["prediction_ml"] / COUNT_DENSITY_DIVISOR).round(2)
    results["prediction_p_strike_annual"] = (
        results["prediction_ng_2025"]
        * results["pln_collection_area_m2"]
        / 1_000_000
        * results["pln_exposure_factor"]
    )
    results["pln_ng_matches_observed_2025"] = np.isclose(
        results["pln_ng"], results["observed_ng_2025"], atol=1e-8
    )
    aligned = results.loc[results["pln_ng_matches_observed_2025"]].copy()
    excluded = results.loc[~results["pln_ng_matches_observed_2025"]].copy()
    if aligned.empty:
        raise DataValidationError("No PLN tower has density aligned with observed 2025 Count")
    top_n = min((TOWER_COUNT + 9) // 10, len(aligned))
    ml_top = set(aligned.nlargest(top_n, "prediction_p_strike_annual")["tower_key"])
    pln_top = set(aligned.nlargest(top_n, "p_strike_annual_pln")["tower_key"])
    rank_correlation = (
        aligned["prediction_p_strike_annual"].rank().corr(aligned["p_strike_annual_pln"].rank())
    )
    by_corridor = {
        corridor: _metrics(group["target_count"], group["prediction_ml"].to_numpy())
        for corridor, group in results.groupby("corridor")
    }
    report = {
        "software_versions": {
            "python": platform.python_version(),
            "numpy": np.__version__,
            "pandas": pd.__version__,
            "scikit_learn": sklearn.__version__,
        },
        "target_definition": "Observed annual lightning Count around each tower, not direct tower strikes/outages",
        "training_target_years": list(TARGET_YEARS[:-1]),
        "selection_folds": selection_folds,
        "final_fit_target_years": list(TARGET_YEARS[:-1]),
        "final_fit_rows": len(final_train),
        "final_test_year": TEST_YEAR,
        "test_protocol_note": "2025 was previously viewed during earlier cleaning and modeling iterations. This is a revised retrospective out-of-time evaluation, not an untouched blind test. Model and feature selection use 2023-2024 only; the selected recipe is refit through 2024 before scoring 2025.",
        "selected_model": selected["model"],
        "selected_feature_group": selected["feature_group"],
        "feature_group_columns": FEATURE_GROUPS,
        "validation_candidates": {"without_grounding": candidates},
        "validation_baseline_last_year": baseline_validation,
        "snapshot_2026_sensitivity": {
            "status": "not_eligible_for_selection_or_historical_prediction",
            "reason": "2026 grounding/protection snapshot has no established 2020-2025 as-of validity; raw Leg D/Total also disagree with Sheet4 for some towers",
            "excluded_raw_columns": [
                "grounding_leg_d_ohm_snapshot_2026",
                "grounding_total_ohm_snapshot_2026",
            ],
            "validation_with_snapshot": sensitivity_metrics,
            "validation_without_snapshot": {
                "mae": next(
                    score["mae"]
                    for score in selected["fold_scores"]
                    if score["validation_target_year"] == 2024
                ),
                "poisson_deviance": next(
                    score["poisson_deviance"]
                    for score in selected["fold_scores"]
                    if score["validation_target_year"] == 2024
                ),
            },
            "source_markers_retained_in_panel": True,
        },
        "test_selected_model": test_metrics,
        "test_baseline_last_year": baseline_test,
        "test_by_corridor": by_corridor,
        "pln_alignment": {
            "metric": "Predicted Count transformed to PLN P_strike_annual scale",
            "count_density_divisor": COUNT_DENSITY_DIVISOR,
            "eligible_towers": len(aligned),
            "excluded_towers": len(excluded),
            "excluded_tower_keys": excluded["tower_key"].tolist(),
            "p_strike_mae": float(
                mean_absolute_error(
                    aligned["p_strike_annual_pln"], aligned["prediction_p_strike_annual"]
                )
            ),
            "spearman_rho": float(rank_correlation),
            "top_62_overlap_count": len(ml_top & pln_top),
            "warning": "The PLN workbook uses 2025 observed lightning density; this is not an independent 2025 forecast. Two source-format fallback values are excluded from alignment, not silently corrected.",
        },
        "data_quality_2020": {
            "reassigned_by_exact_gps": int(matched_2020.sum()),
            "unresolved_missing_count": int(unresolved_2020.sum()),
            "usage": "2020 Count follows exact GPS match by user-directed assumption; unmatched values stay unassigned and missing lag values are imputed from train only",
            "raw_source_modified": False,
        },
    }
    output_dir.mkdir(parents=True, exist_ok=True)
    report_temp = report_path.with_suffix(".tmp.json")
    predictions_temp = predictions_path.with_suffix(".tmp.csv")
    try:
        report_temp.write_text(
            json.dumps(report, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
        )
        results.sort_values(["corridor", "tower_key"]).to_csv(predictions_temp, index=False)
        report_temp.replace(report_path)
        predictions_temp.replace(predictions_path)
    finally:
        for temporary in (report_temp, predictions_temp):
            if temporary.exists():
                temporary.unlink()
    return EvaluationResult(report_path=report_path, predictions_path=predictions_path)
