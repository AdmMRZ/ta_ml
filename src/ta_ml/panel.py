"""Build the leakage-aware next-year model panel from the clean source workbook."""

from __future__ import annotations

import csv
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime
from math import isfinite
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .cleaning import DataValidationError
from .config import DEVELOPMENT_END_YEAR, FIRST_OBSERVED_YEAR, TARGET_YEARS, TEST_YEAR, TOWER_COUNT


@dataclass(frozen=True)
class PanelSummary:
    output_path: Path
    rows: int


def _table(workbook: Any, name: str) -> list[dict[str, Any]]:
    worksheet = workbook[name]
    iterator = worksheet.iter_rows(values_only=True)
    headers = next(iterator)
    return [dict(zip(headers, row, strict=True)) for row in iterator]


def _numeric_count(value: object, label: str) -> int:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise DataValidationError(f"{label} is not a numeric count: {value!r}")
    if not isfinite(value) or value < 0 or not float(value).is_integer():
        raise DataValidationError(f"{label} is not a non-negative integer: {value!r}")
    return int(value)


def _tla_asof(registry: dict[str, Any], feature_year: int) -> int | None:
    install_date = registry["tla_install_date"]
    tla_count = registry["tla_count"]
    if install_date is None:
        if isinstance(tla_count, (int, float)) and tla_count > 0:
            return None
        return 0
    if not isinstance(install_date, datetime):
        raise DataValidationError(f"Unparseable TLA installation date: {install_date!r}")
    return int(install_date.year <= feature_year)


def build_model_panel(
    *, clean_path: Path, output_path: Path, overwrite: bool = False
) -> PanelSummary:
    """Create next-year pairs; never interpret 2026 snapshots as historical facts."""

    clean_path = clean_path.resolve()
    output_path = output_path.resolve()
    if not clean_path.is_file():
        raise DataValidationError(f"Clean workbook does not exist: {clean_path}")
    if clean_path == output_path:
        raise DataValidationError("Panel output cannot replace the clean workbook")
    if output_path.exists() and not overwrite:
        raise FileExistsError(f"Panel output exists: {output_path}")

    workbook = load_workbook(clean_path, read_only=True, data_only=True)
    try:
        historical = _table(workbook, "tower_year")
        registry_rows = _table(workbook, "tower_registry")
        coordinate_audit = _table(workbook, "coordinate_2020_audit")
    finally:
        workbook.close()

    registry = {(r["corridor"], r["tower_key"]): r for r in registry_rows}
    if len(registry) != TOWER_COUNT or len(registry_rows) != TOWER_COUNT:
        raise DataValidationError("Registry must contain exactly 617 unique corridor/tower keys")

    by_tower: dict[tuple[str, str], dict[int, dict[str, Any]]] = defaultdict(dict)
    for row in historical:
        key = (row["corridor"], row["tower_key"])
        year = row["year"]
        if key not in registry or year in by_tower[key]:
            raise DataValidationError(f"Unknown or duplicate historical key: {key}, {year}")
        by_tower[key][year] = row

    if set(by_tower) != set(registry):
        raise DataValidationError("Historical panel and registry have different tower keys")

    coordinate_candidates: dict[tuple[str, str], dict[str, Any]] = {}
    if len(coordinate_audit) != 617:
        raise DataValidationError("2020 coordinate audit must cover all 617 source rows")
    for audit_row in coordinate_audit:
        if audit_row["match_status"] != "exact_coordinate":
            continue
        matched_key = (audit_row["matched_corridor"], audit_row["matched_tower_key"])
        if matched_key not in registry or matched_key in coordinate_candidates:
            raise DataValidationError(f"Invalid/duplicate 2020 coordinate match: {matched_key}")
        coordinate_candidates[matched_key] = audit_row

    panel: list[dict[str, Any]] = []
    for key in sorted(registry):
        years = by_tower[key]
        if set(years) != set(range(FIRST_OBSERVED_YEAR, TEST_YEAR + 1)):
            raise DataValidationError(f"Incomplete history for {key}: {sorted(years)}")
        static = registry[key]
        registry_coords = (static["latitude_2026"], static["longitude_2026"])
        # 2021-2025 coordinates are independently present in the lightning source.
        # Every one must agree with the named tower in the 2026 registry.
        for year in TARGET_YEARS:
            observed = (years[year]["latitude"], years[year]["longitude"])
            if observed != registry_coords:
                raise DataValidationError(f"Coordinate mismatch for {key}, {year}")

        for target_year in TARGET_YEARS:
            feature_year = target_year - 1
            coordinate_candidate = coordinate_candidates.get(key) if target_year == 2021 else None
            previous_counts = {
                year: (
                    _numeric_count(years[year]["count"], f"{key}, {year}")
                    if years[year]["count"] is not None
                    else None
                )
                for year in range(FIRST_OBSERVED_YEAR, feature_year + 1)
            }
            past = [value for value in previous_counts.values() if value is not None]
            lag1 = previous_counts.get(feature_year)
            lag2 = previous_counts.get(feature_year - 1)
            lag3 = previous_counts.get(feature_year - 2)
            if target_year == 2021:
                candidate_count = (
                    coordinate_candidate["source_count_2020"] if coordinate_candidate else None
                )
                if candidate_count != lag1:
                    raise DataValidationError(f"2020 reconciliation disagrees with audit: {key}")
            row: dict[str, Any] = {
                "corridor": key[0],
                "tower_key": key[1],
                "tower_name": static["tower_name"],
                "feature_year": feature_year,
                "target_year": target_year,
                "split": "development" if target_year <= DEVELOPMENT_END_YEAR else "test",
                "target_count": _numeric_count(
                    years[target_year]["count"], f"{key}, {target_year}"
                ),
                "count_lag1": lag1,
                "count_lag2": lag2,
                "count_lag3": lag3,
                "count_mean_available": sum(past) / len(past) if past else None,
                "count_history_years": len(past),
                "count_2020_unresolved": int(target_year == 2021 and coordinate_candidate is None),
                "count_2020_reconciled_source_row": (
                    coordinate_candidate["source_row"] if coordinate_candidate else None
                ),
                "count_2020_coordinate_match_status": (
                    "exact_coordinate_reassigned"
                    if coordinate_candidate
                    else "unmatched"
                    if target_year == 2021
                    else "not_applicable"
                ),
                "latitude_verified": static["latitude_2026"],
                "longitude_verified": static["longitude_2026"],
                "elevation_m": static["elevation_m"],
                "tower_height_m": static["tower_height_m"],
                "span_left_m": static["span_left_m"],
                "span_right_m": static["span_right_m"],
                "avg_neighbor_distance_m": static["avg_neighbor_distance_m"],
                "terrain_factor": static["terrain_factor"],
                "angle_exposure_factor": static["angle_exposure_factor"],
                "tla_active_asof": _tla_asof(static, feature_year),
                "tla_install_date_source": static["tla_install_date"],
                "grounding_leg_a_ohm_snapshot_2026": static["grounding_leg_a_ohm"],
                "grounding_leg_b_ohm_snapshot_2026": static["grounding_leg_b_ohm"],
                "grounding_leg_c_ohm_snapshot_2026": static["grounding_leg_c_ohm"],
                "grounding_leg_d_ohm_snapshot_2026": static["grounding_leg_d_ohm"],
                "grounding_total_ohm_snapshot_2026": static["grounding_total_source_ohm"],
                "grounding_leg_d_sheet4_reference_ohm": static[
                    "grounding_leg_d_sheet4_reference_ohm"
                ],
                "grounding_total_sheet4_reference_ohm": static[
                    "grounding_total_sheet4_reference_ohm"
                ],
                "grounding_d_matches_sheet4": static["grounding_d_matches_sheet4"],
                "grounding_total_matches_sheet4": static["grounding_total_matches_sheet4"],
                "grounding_rod_count_snapshot_2026": static["grounding_rod_count"],
                "soil_resistivity_snapshot_2026": static["soil_resistivity_source"],
                "protection_angle_deg_snapshot_2026": static["protection_angle_deg"],
                "insulator_type_snapshot_2026": static["insulator_type"],
                "circuit_configuration_snapshot_2026": static["circuit_configuration"],
            }
            panel.append(row)

    if len(panel) != TOWER_COUNT * len(TARGET_YEARS) or Counter(row["split"] for row in panel) != {
        "development": TOWER_COUNT * (len(TARGET_YEARS) - 1),
        "test": TOWER_COUNT,
    }:
        raise DataValidationError("Unexpected model panel population or temporal split")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = output_path.with_suffix(".tmp.csv")
    try:
        with temporary_path.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.DictWriter(handle, fieldnames=list(panel[0]))
            writer.writeheader()
            writer.writerows(panel)
        temporary_path.replace(output_path)
    finally:
        if temporary_path.exists():
            temporary_path.unlink()
    return PanelSummary(output_path=output_path, rows=len(panel))
